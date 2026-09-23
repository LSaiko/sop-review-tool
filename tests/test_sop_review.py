"""Offline pytest suite for sop_review.py (21 CFR 820 SOP gap analysis).

No test makes a live Anthropic API call: ``anthropic.Anthropic`` is replaced with a fake
client, so the suite runs with no ANTHROPIC_API_KEY and no network. Live-model
re-validation stays in ``tests/revalidate.py``.
"""

import json
import sys
from pathlib import Path
from types import SimpleNamespace

import anthropic
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import sop_review as sr  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent

GOOD_REVIEW = {
    "overall_status": "NEEDS REVISION",
    "overall_rationale": "Two 21 CFR 820 items incomplete.",
    "findings": [
        {"id": "U01", "item": "Purpose Statement", "regulation": "21 CFR 820.40",
         "status": "PRESENT", "evidence": "1.0 PURPOSE", "recommendation": None},
        {"id": "C01", "item": "Cleaning agents", "regulation": "21 CFR 820.70(e)",
         "status": "INCOMPLETE", "evidence": "detergent", "recommendation": "Name the agent."},
    ],
}


class FakeClient:
    """Stands in for anthropic.Anthropic; records calls, returns queued responses/errors."""

    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        self.calls.append(kwargs)
        r = self.responses.pop(0)
        if isinstance(r, Exception):
            raise r
        return SimpleNamespace(content=[SimpleNamespace(text=r)])


@pytest.fixture
def fake_claude(monkeypatch):
    def install(*responses):
        client = FakeClient(responses)
        monkeypatch.setattr(sr.anthropic, "Anthropic", lambda *a, **k: client)
        monkeypatch.setattr(sr.time, "sleep", lambda s: None)
        return client
    return install


# --- checklist ----------------------------------------------------------------------------

def test_checklist_has_universal_and_type_items():
    txt = sr.build_checklist_text("inspection")
    assert "[U01]" in txt and "[U11]" in txt and "[I01]" in txt


def test_unknown_type_gets_universal_only():
    txt = sr.build_checklist_text("does-not-exist")
    assert "[U01]" in txt and "[I01]" not in txt


# --- response parsing ---------------------------------------------------------------------

def test_fenced_json_parses():
    parsed = sr.parse_claude_response("```json\n" + json.dumps(GOOD_REVIEW) + "\n```")
    assert parsed["overall_status"] == "NEEDS REVISION"


@pytest.mark.parametrize("raw", [
    "not json",
    json.dumps({"overall_status": "X", "findings": []}),                 # missing keys
    json.dumps({**GOOD_REVIEW, "findings": []}),                          # empty findings
    json.dumps({**GOOD_REVIEW, "findings": [{"id": "U01", "status": "MAYBE"}]}),
    json.dumps({**GOOD_REVIEW, "findings": [{"id": "U01"}]}),             # no status
])
def test_malformed_response_rejected(raw):
    with pytest.raises(ValueError):
        sr.parse_claude_response(raw)


# --- SOP file reading ---------------------------------------------------------------------

def test_missing_file():
    with pytest.raises(FileNotFoundError):
        sr.read_sop_file("no_such_file_12345.txt")


def test_unsupported_extension(tmp_path):
    p = tmp_path / "x.rtf"
    p.write_text("hi", encoding="utf-8")
    with pytest.raises(ValueError):
        sr.read_sop_file(str(p))


def test_reads_txt(tmp_path):
    p = tmp_path / "ok.txt"
    p.write_text("hello sop", encoding="utf-8")
    assert sr.read_sop_file(str(p)) == "hello sop"


def test_non_utf8_txt_rejected(tmp_path):
    p = tmp_path / "bad.txt"
    p.write_bytes(b"\xff\xfe\x00bad bytes")
    with pytest.raises(ValueError):
        sr.read_sop_file(str(p))


def test_reads_docx_paragraphs_and_tables(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("1.0 PURPOSE")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "Torque 5 Nm"
    p = tmp_path / "sop.docx"
    doc.save(p)
    text = sr.read_sop_file(str(p))
    assert "1.0 PURPOSE" in text and "Torque 5 Nm" in text


def test_reads_pdf(tmp_path):
    # The bundled example report is a real PDF with extractable text.
    assert "21 CFR" in sr.read_sop_file(str(ROOT / "sop_review_report.pdf"))


# --- PDF rendering ------------------------------------------------------------------------

def test_pdf_escapes_markup_chars(tmp_path):
    nasty = "Tolerance < 5 µm & spec > 3 <b>bold?</b>"
    review = {"overall_status": "NEEDS REVISION", "overall_rationale": nasty, "findings": [{
        "id": "U01", "item": "Purpose & Scope <x>", "regulation": "21 CFR 820.40 <a>",
        "status": "INCOMPLETE", "evidence": nasty, "recommendation": nasty}]}
    out = tmp_path / "report.pdf"
    sr.generate_pdf_report(review, str(out), "sample.txt", "III", "inspection")
    assert out.stat().st_size > 0


def test_scorecard_empty_findings_no_zero_division():
    sr._scorecard_table([], sr._make_styles())


# --- Claude call (mocked) -----------------------------------------------------------------

def test_call_claude_sends_sop_and_checklist(fake_claude):
    client = fake_claude(json.dumps(GOOD_REVIEW))
    raw = sr.call_claude("MY SOP TEXT", "cleaning", "II")
    assert json.loads(raw) == GOOD_REVIEW
    prompt = client.calls[0]["messages"][0]["content"]
    assert "MY SOP TEXT" in prompt and "DEVICE CLASS: II" in prompt and "[U01]" in prompt


def _api_error():
    return anthropic.APIConnectionError(request=SimpleNamespace())


def test_call_claude_retries_then_succeeds(fake_claude):
    client = fake_claude(_api_error(), json.dumps(GOOD_REVIEW))
    assert json.loads(sr.call_claude("x", "cleaning", "II")) == GOOD_REVIEW
    assert len(client.calls) == 2


def test_call_claude_gives_up_after_three(fake_claude):
    fake_claude(_api_error(), _api_error(), _api_error())
    with pytest.raises(anthropic.APIError):
        sr.call_claude("x", "cleaning", "II")


# --- CLI end to end (mocked) --------------------------------------------------------------

def _run_main(monkeypatch, *argv):
    monkeypatch.setattr(sys, "argv", ["sop_review.py", *argv])
    sr.main()


def test_main_without_api_key_exits(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    with pytest.raises(SystemExit) as e:
        _run_main(monkeypatch, "--file", str(ROOT / "sample_sop.txt"),
                  "--sop-type", "cleaning", "--device-class", "II")
    assert e.value.code == 1


def test_main_end_to_end(monkeypatch, tmp_path, fake_claude):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-not-a-real-key")
    fake_claude(json.dumps(GOOD_REVIEW))
    pdf, js = tmp_path / "r.pdf", tmp_path / "r.json"
    _run_main(monkeypatch, "--file", str(ROOT / "sample_sop.txt"), "--sop-type", "cleaning",
              "--device-class", "II", "--output", str(pdf), "--json-output", str(js))
    assert pdf.stat().st_size > 0
    assert json.loads(js.read_text(encoding="utf-8")) == GOOD_REVIEW


def test_main_bad_claude_output_exits(monkeypatch, tmp_path, fake_claude):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-not-a-real-key")
    fake_claude("not json")
    with pytest.raises(SystemExit) as e:
        _run_main(monkeypatch, "--file", str(ROOT / "sample_sop.txt"), "--sop-type",
                  "cleaning", "--device-class", "II", "--output", str(tmp_path / "r.pdf"))
    assert e.value.code == 1


def test_no_hardcoded_api_key_in_source():
    src = (ROOT / "sop_review.py").read_text(encoding="utf-8")
    assert 'os.getenv("ANTHROPIC_API_KEY")' in src
    # The only sk-ant string allowed is the placeholder in the help message.
    assert src.count("sk-ant-") == src.count("sk-ant-...")


# --- Part 11 hook (optional; HTTP layer faked) --------------------------------------------

import part11  # noqa: E402

MAIN_ARGS = ("--file", str(ROOT / "sample_sop.txt"), "--sop-type", "cleaning",
             "--device-class", "II")


@pytest.fixture
def fake_part11(monkeypatch):
    calls = []

    def fake_request(method, path, body=None, headers=None):
        calls.append((method, path, body, headers))
        if path.startswith("/records/"):
            return {"current_hash": "a" * 64}
        return {"id": "evt-1"}
    monkeypatch.setattr(part11, "_request", fake_request)
    return calls


def test_main_unchanged_when_part11_url_unset(monkeypatch, tmp_path, fake_claude, fake_part11):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-not-a-real-key")
    monkeypatch.delenv("PART11_AUDIT_URL", raising=False)
    fake_claude(json.dumps(GOOD_REVIEW))
    _run_main(monkeypatch, *MAIN_ARGS, "--output", str(tmp_path / "r.pdf"))
    assert fake_part11 == []


def test_main_records_review_when_part11_url_set(monkeypatch, tmp_path, fake_claude,
                                                  fake_part11, capsys):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-not-a-real-key")
    monkeypatch.setenv("PART11_AUDIT_URL", "http://part11.invalid/")
    monkeypatch.setenv("PART11_AUDIT_TOKEN", "tok")
    fake_claude(json.dumps(GOOD_REVIEW))
    _run_main(monkeypatch, *MAIN_ARGS, "--output", str(tmp_path / "r.pdf"))
    (m1, p1, b1, _), (m2, p2, b2, h2) = fake_part11
    assert (m1, p1.split("/")[1:3]) == ("PUT", ["records", "sop_review"])
    assert b1["content"]["review"] == GOOD_REVIEW
    record_id = p1.rsplit("/", 1)[1]
    assert (m2, p2) == ("POST", "/events")
    assert b2 == {"actor": b1["actor"], "action": "sop_review:completed",
                  "record_type": "sop_review", "record_id": record_id, "after_hash": "a" * 64}
    assert h2 == {"X-Audit-Token": "tok"}
    assert f"--record-id {record_id}" in capsys.readouterr().out


def test_part11_down_does_not_fail_review(monkeypatch, tmp_path, fake_claude, capsys):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-not-a-real-key")
    monkeypatch.setenv("PART11_AUDIT_URL", "http://127.0.0.1:9")  # nothing listens here
    fake_claude(json.dumps(GOOD_REVIEW))
    pdf = tmp_path / "r.pdf"
    _run_main(monkeypatch, *MAIN_ARGS, "--output", str(pdf))
    assert pdf.stat().st_size > 0
    assert "Part 11 audit trail not updated" in capsys.readouterr().err


def test_sign_approval_posts_approved_meaning(fake_part11):
    part11.sign_approval("SOP-1", "alice", "pw")
    assert fake_part11 == [("POST", "/sign", {"username": "alice", "password": "pw",
                                              "record_id": "SOP-1", "meaning": "approved"}, None)]


def test_part11_cli_requires_url(monkeypatch):
    monkeypatch.delenv("PART11_AUDIT_URL", raising=False)
    monkeypatch.setattr(sys, "argv", ["part11.py", "verify", "--signature-id", "s1"])
    with pytest.raises(SystemExit, match="PART11_AUDIT_URL"):
        part11.main()


def test_part11_cli_sign_prompts_for_password(monkeypatch, fake_part11):
    monkeypatch.setenv("PART11_AUDIT_URL", "http://part11.invalid")
    monkeypatch.setattr(part11.getpass, "getpass", lambda prompt: "pw")
    monkeypatch.setattr(sys, "argv", ["part11.py", "sign", "--record-id", "R1", "--signer", "alice"])
    part11.main()
    assert fake_part11[0][2]["password"] == "pw"


def test_part11_cli_verify(monkeypatch, fake_part11):
    monkeypatch.setenv("PART11_AUDIT_URL", "http://part11.invalid")
    monkeypatch.setattr(sys, "argv", ["part11.py", "verify", "--signature-id", "s1"])
    part11.main()
    assert fake_part11 == [("GET", "/verify-signature/s1", None, None)]
